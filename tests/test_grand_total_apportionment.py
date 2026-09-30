"""The grand total as arithmetic, not just a cross-check.

The circled Grand Total is the biggest, clearest, most deliberate figure on
these tickets. When exactly one line has no price, it gives that price exactly.

The residual was already being computed — for a debug-trace message — and then
discarded, while the blank cell went red and the reviewer was asked to work out
a number the tool was holding. These tests pin both halves: it fills the line it
can prove, and it refuses every case where the arithmetic doesn't actually
determine an answer.
"""
import io

from openpyxl import load_workbook

from app.db import db
from app.pipeline.assemble import assemble_and_persist
from app.sheets.write import write_review_workbook

AMBER = "FFF2CC"


def _f(value, confidence="high"):
    return {"value": value, "confidence": confidence}


def _empty_label():
    return {"gtin": None, "lot": None, "expiry": None, "mfg": None,
            "serial": None, "raw": None, "decoded": False, "ref": None}


def _run(prices, grand_total, *, freight=None, qtys=None, wasted=None):
    """prices: per-line unit price, None for an unpriced line."""
    batch = db.create_batch()
    t = db.create_ticket({"batch_id": batch["id"], "source_filename": "MO-gt.jpg",
                          "status": "pending_review"})
    lines = []
    for i, p in enumerate(prices):
        line = {"index": i, "ref": _f(f"GT-REF-{i}"),
                "qty": _f((qtys or {}).get(i)),
                "unit_price": _f(p, "high" if p is not None else "low")}
        if (wasted or {}).get(i):
            line["wasted"] = _f(True)
        lines.append(line)
    summary = assemble_and_persist(db.get_ticket(t["ticket_id"]), {
        "header": {"surgery_date": _f("2026-06-01")},
        "lines": lines,
        "freight": _f(freight, "high" if freight is not None else "low"),
        "grand_total": _f(grand_total),
    }, [_empty_label() for _ in prices])
    rows = sorted(db.lines_for_ticket(t["ticket_id"]),
                  key=lambda x: x.get("created_at") or "")
    return batch, rows, summary


def _price_cells(batch):
    ws = load_workbook(io.BytesIO(write_review_workbook(batch["id"])))["Usage"]
    headers = [c.value for c in ws[1]]
    col = headers.index("Price") + 1
    return [ws.cell(row=r, column=col) for r in range(2, ws.max_row + 1)]


def _fill(cell) -> str:
    if cell.fill is None or cell.fill.patternType is None:
        return ""
    return str(cell.fill.start_color.rgb or "")


# ---------------------------------------------------------------------------
# The case it can prove
# ---------------------------------------------------------------------------
def test_one_blank_line_is_filled_from_the_residual():
    """Cox ticket, in miniature: 1300 + 1150 + 900 + 250 against a circled 3775
    leaves 175, and it belongs to the one line without a price."""
    _, rows, summary = _run([1300, 1150, 900, 250, None], 3775)
    assert rows[4]["unit_price"] == 175.0
    assert rows[4]["line_total"] == 175.0
    assert not any("Grand total" in f for f in summary["flags"]), \
        "the ticket now reconciles, so the mismatch flag should be gone"


def test_the_derived_price_is_amber_and_says_where_it_came_from():
    batch, rows, _ = _run([1300, 1150, 900, 250, None], 3775)
    cell = _price_cells(batch)[4]
    assert cell.value == 175.0
    assert _fill(cell).endswith(AMBER)
    assert "derived from the grand total" in " ".join(rows[4]["flags"])


def test_a_derived_price_is_never_promoted_to_confident():
    """After filling, the lines add up to the total BY CONSTRUCTION. Treating
    that as independent agreement would be the tool citing its own arithmetic."""
    batch, _, _ = _run([1300, 1150, 900, 250, None], 3775)
    for cell in _price_cells(batch):
        assert _fill(cell).endswith(AMBER), \
            "no price may go white on the strength of a total we just satisfied"


def test_freight_is_subtracted_before_apportioning():
    _, rows, _ = _run([1000, None], 1250, freight=50)
    assert rows[1]["unit_price"] == 200.0


def test_an_unread_freight_is_called_out():
    """The residual silently absorbs a delivery fee nobody read — the single
    most likely way this rule produces a wrong number."""
    _, rows, summary = _run([1000, None], 1250)
    assert rows[1]["unit_price"] == 250.0
    assert "freight" in " ".join(rows[1]["flags"]).lower()
    assert any("freight" in f.lower() for f in summary["flags"])


def test_quantity_divides_the_residual():
    _, rows, _ = _run([1000, None], 1400, qtys={1: 4})
    assert rows[1]["unit_price"] == 100.0
    assert rows[1]["line_total"] == 400.0


def test_a_residual_that_does_not_divide_evenly_is_filled_and_noted():
    _, rows, _ = _run([1000, None], 1100, qtys={1: 3})
    assert rows[1]["unit_price"] == 33.33
    assert "does not divide evenly" in " ".join(rows[1]["flags"])


# ---------------------------------------------------------------------------
# The cases it must refuse
# ---------------------------------------------------------------------------
def test_two_blank_lines_are_not_split():
    _, rows, summary = _run([1000, None, None], 1500)
    assert rows[1]["unit_price"] is None and rows[2]["unit_price"] is None
    joined = " ".join(summary["flags"])
    assert "$500.00 unaccounted" in joined and "2 lines" in joined
    assert "250.00 each" in joined, "the per-line average is still worth saying"


def test_a_negative_residual_writes_nothing():
    _, rows, summary = _run([1000, None], 800)
    assert rows[1]["unit_price"] is None
    assert any("Grand total" in f for f in summary["flags"])


def test_a_zero_residual_never_writes_a_zero_price():
    """A zero price asserts "this was free", which is a different claim from
    "we don't know"."""
    _, rows, _ = _run([1000, None], 1000)
    assert rows[1]["unit_price"] is None
    assert "may not be billable" in " ".join(rows[1]["flags"])


def test_an_implausible_residual_is_reported_rather_than_filled():
    _, rows, _ = _run([100, None], 20000)
    assert rows[1]["unit_price"] is None
    assert "too large to be plausible" in " ".join(rows[1]["flags"])


def test_a_wasted_line_is_never_apportioned():
    """Whether a wasted component is billed is a business decision."""
    _, rows, _ = _run([1000, None], 1200, wasted={1: True})
    assert rows[1]["unit_price"] is None
    assert "business decision" in " ".join(rows[1]["flags"])


def test_no_grand_total_means_no_apportionment():
    _, rows, _ = _run([1000, None], None)
    assert rows[1]["unit_price"] is None


def test_a_fully_priced_ticket_is_untouched_and_still_promotes():
    """The independent-agreement promotion must survive: when nothing was
    derived, lines summing to the total is still real evidence."""
    batch, _, summary = _run([1000, 500], 1500)
    assert not any("Grand total" in f for f in summary["flags"])
    for cell in _price_cells(batch):
        assert _fill(cell) == "", "an independently reconciled price is confident"
