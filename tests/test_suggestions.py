"""Red means "nothing to suggest", not "we gave up".

A batch of nine real tickets came back with 54% of the deliverable red and the
run reported success. Most of that was one cause — the vision call returned
nothing on every ticket, silently (see test_vision_failure.py). But it exposed a
second, independent problem worth fixing on its own: when the tool DID hold a
candidate value, it threw it away and showed an empty red cell anyway.

A sub-threshold read was deleted. The hospital written on the ticket never
reached the Usage sheet. The number the ticket's own grand total pins down
exactly was computed for a log message and discarded. In every one of those
cases the reviewer was asked to find something the tool was already holding.

These tests pin the new rule: an amber cell always carries a value and says in
Notes where it came from; an empty cell is always red.
"""
import io

from openpyxl import load_workbook

from app.db import db
from app.pipeline.assemble import assemble_and_persist
from app.sheets.write import write_review_workbook

AMBER, RED = "FFF2CC", "F4CCCC"


def _f(value, confidence="high"):
    return {"value": value, "confidence": confidence}


def _empty_label():
    return {"gtin": None, "lot": None, "expiry": None, "mfg": None,
            "serial": None, "raw": None, "decoded": False, "ref": None}


def _ticket(**over):
    batch = db.create_batch()
    row = {"batch_id": batch["id"], "source_filename": "MO-sugg.jpg",
           "status": "pending_review"}
    row.update(over)
    t = db.create_ticket(row)
    return batch, db.get_ticket(t["ticket_id"])


def _usage(batch):
    ws = load_workbook(io.BytesIO(write_review_workbook(batch["id"])))["Usage"]
    headers = [c.value for c in ws[1]]
    return ws, headers


def _cell(batch, header, row=2):
    ws, headers = _usage(batch)
    return ws.cell(row=row, column=headers.index(header) + 1)


def _fill(cell) -> str:
    if cell.fill is None or cell.fill.patternType is None:
        return ""
    return str(cell.fill.start_color.rgb or "")


def _run(batch, ticket, header=None, lines=None, grand_total=None, freight=None):
    assemble_and_persist(ticket, {
        "header": header or {},
        "lines": lines if lines is not None else [
            {"index": 0, "ref": _f("SUGG-REF-1"), "qty": _f(1), "unit_price": _f(100)}],
        "freight": _f(freight, "high" if freight is not None else "low"),
        "grand_total": _f(grand_total, "high" if grand_total is not None else "low"),
    }, [_empty_label() for _ in range(len(lines) if lines is not None else 1)])


# ---------------------------------------------------------------------------
# A read the model was unsure of is shown, not deleted
# ---------------------------------------------------------------------------
def test_a_sub_threshold_read_is_shown_amber_instead_of_deleted():
    batch, ticket = _ticket()
    _run(batch, ticket, header={"surgeon": _f("Woodworth", "low"),
                                "surgery_date": _f("2026-06-01")})
    cell = _cell(batch, "Surgeon")
    assert cell.value == "Woodworth", "the read was deleted instead of offered"
    assert _fill(cell).endswith(AMBER)


def test_a_field_that_was_never_read_stays_red_and_blank():
    batch, ticket = _ticket()
    _run(batch, ticket, header={"surgeon": _f(None, "low"),
                                "surgery_date": _f("2026-06-01")})
    cell = _cell(batch, "Surgeon")
    assert cell.value is None
    assert _fill(cell).endswith(RED)


def test_the_ticket_says_which_fields_the_model_was_unsure_of():
    """Bare amber means an ordinary read. Amber with a reason in Notes is the
    difference between "check this" and "check this, here's why"."""
    batch, ticket = _ticket()
    _run(batch, ticket, header={"surgeon": _f("Woodworth", "low"),
                                "surgery_date": _f("2026-06-01")})
    notes = _cell(batch, "Notes").value or ""
    assert "unsure" in notes.lower() and "Surgeon" in notes


def test_an_unparseable_date_renders_red_not_empty_amber():
    """A confidence attached to something that formats to nothing would leave an
    empty amber cell, which asks the reviewer to check a blank."""
    batch, ticket = _ticket()
    _run(batch, ticket, header={"surgery_date": _f("not a date")})
    cell = _cell(batch, "Date")
    assert cell.value is None
    assert _fill(cell).endswith(RED)


# ---------------------------------------------------------------------------
# Hospital: four rungs, best evidence first
# ---------------------------------------------------------------------------
def test_the_handwritten_hospital_reaches_the_usage_sheet():
    """The Usage column used to be fed only by the surgeon lookup, so a ticket
    that plainly read ENLOE came back with an empty red Hospital cell."""
    batch, ticket = _ticket()
    _run(batch, ticket, header={"hospital": _f("ENLOE"),
                                "surgery_date": _f("2026-06-01")})
    cell = _cell(batch, "Hospital")
    assert cell.value == "ENLOE"
    assert _fill(cell).endswith(AMBER)


def test_the_handwritten_hospital_and_the_surgeon_record_agreeing_is_confident():
    db.replace_reference_surgeons([
        {"surgeon_distcode": "KONKINRO-MO-001", "surgeon_last_name": "Konkin",
         "dist_code": "RO-MO-001", "status": "Active",
         "surgeon_full_name": "K Konkin", "hospital": "Enloe", "region": "W",
         "distributor_rep": "R"}])
    batch, ticket = _ticket()
    _run(batch, ticket, header={"hospital": _f("ENLOE"), "surgeon": _f("Konkin"),
                                "rep_code": _f("RO-MO-001"),
                                "surgery_date": _f("2026-06-01")})
    cell = _cell(batch, "Hospital")
    assert cell.value == "ENLOE"
    assert _fill(cell) == "", "two independent sources agreed — no fill"


def test_a_disagreement_prefers_the_ticket_and_says_so():
    """What the surgeon wrote on THIS ticket outranks what the master said when
    it was last exported — but the reviewer is told they disagree."""
    db.replace_reference_surgeons([
        {"surgeon_distcode": "CHASEJT-MO-001", "surgeon_last_name": "Chase",
         "dist_code": "JT-MO-001", "status": "Active", "surgeon_full_name": "J Chase",
         "hospital": "Old General", "region": "W", "distributor_rep": "R"}])
    batch, ticket = _ticket()
    _run(batch, ticket, header={"hospital": _f("Paragon Surgical Center"),
                                "surgeon": _f("Chase"), "rep_code": _f("JT-MO-001"),
                                "surgery_date": _f("2026-06-01")})
    cell = _cell(batch, "Hospital")
    assert cell.value == "Paragon Surgical Center"
    assert _fill(cell).endswith(AMBER)
    assert "Old General" in (_cell(batch, "Notes").value or "")


def test_no_handwritten_hospital_falls_back_to_the_surgeon_record():
    db.replace_reference_surgeons([
        {"surgeon_distcode": "MILLERJS-MO-001", "surgeon_last_name": "Miller",
         "dist_code": "JS-MO-001", "status": "Active", "surgeon_full_name": "M Miller",
         "hospital": "Health First Viera", "region": "S", "distributor_rep": "R"}])
    batch, ticket = _ticket()
    _run(batch, ticket, header={"surgeon": _f("Miller"), "rep_code": _f("JS-MO-001"),
                                "surgery_date": _f("2026-06-01")})
    assert _cell(batch, "Hospital").value == "Health First Viera"


def test_an_unread_surgeon_can_still_yield_a_hospital_from_the_distcode():
    """The surgeon chain needs BOTH a name and a DistCode, so a missed surgeon
    read used to take the hospital down with it — even when the code was clean
    and had only ever belonged to one account."""
    db.replace_reference_surgeons([
        {"surgeon_distcode": "AAA1ZZ-MO-009", "surgeon_last_name": "Aaa",
         "dist_code": "ZZ-MO-009", "status": "Active", "surgeon_full_name": "A Aaa",
         "hospital": "Solo Regional", "region": "S", "distributor_rep": "R"},
        {"surgeon_distcode": "BBB1ZZ-MO-009", "surgeon_last_name": "Bbb",
         "dist_code": "ZZ-MO-009", "status": "Active", "surgeon_full_name": "B Bbb",
         "hospital": "Solo Regional", "region": "S", "distributor_rep": "R"}])
    batch, ticket = _ticket()
    _run(batch, ticket, header={"rep_code": _f("ZZ-MO-009"),
                                "surgery_date": _f("2026-06-01")})
    cell = _cell(batch, "Hospital")
    assert cell.value == "Solo Regional"
    assert _fill(cell).endswith(AMBER)
    assert "DistCode" in (_cell(batch, "Notes").value or "")


def test_an_ambiguous_distcode_offers_nothing():
    """Two genuinely different accounts is not a suggestion, it's a coin toss."""
    db.replace_reference_surgeons([
        {"surgeon_distcode": "AAA1QQ-MO-009", "surgeon_last_name": "Aaa",
         "dist_code": "QQ-MO-009", "status": "Active", "surgeon_full_name": "A Aaa",
         "hospital": "North Regional", "region": "S", "distributor_rep": "R"},
        {"surgeon_distcode": "BBB1QQ-MO-009", "surgeon_last_name": "Bbb",
         "dist_code": "QQ-MO-009", "status": "Active", "surgeon_full_name": "B Bbb",
         "hospital": "South Medical Center", "region": "S", "distributor_rep": "R"}])
    batch, ticket = _ticket()
    _run(batch, ticket, header={"rep_code": _f("QQ-MO-009"),
                                "surgery_date": _f("2026-06-01")})
    cell = _cell(batch, "Hospital")
    assert cell.value is None
    assert _fill(cell).endswith(RED)


def test_usage_and_tickets_agree_about_the_hospital():
    """They were resolved by two separate code paths and could differ."""
    batch, ticket = _ticket()
    _run(batch, ticket, header={"hospital": _f("Schneider Hospital"),
                                "surgery_date": _f("2026-06-01")})
    wb = load_workbook(io.BytesIO(write_review_workbook(batch["id"])))
    tk = wb["Tickets"]
    th = [c.value for c in tk[1]]
    assert (tk.cell(row=2, column=th.index("Hospital") + 1).value
            == _cell(batch, "Hospital").value == "Schneider Hospital")
