"""Patient-initials extraction (app/pipeline/patient.py) and the PHI guarantees
that have to survive it.

The feature is a deliberate, flag-gated exception to "we do not process patient
data at all". These tests pin the boundaries of that exception: off by default,
never runs when the redaction gate fails, only the sticker crop is ever sent,
and the image that gets STORED is still the fully redacted one.
"""
import io
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import cv2
import numpy as np
from openpyxl import load_workbook

import app.pipeline.run as run
from app.db import db
from app.pipeline import patient
from app.pipeline.assemble import assemble_and_persist
from app.pipeline.run import ingest_image
from app.pipeline.template import MAXX_ORTHO, geometry_for
from app.sheets.write import write_review_workbook

# Maxx Orthopedics patient sticker on a 600x400 frame: x 330..582, y 20..84.
_W, _H = 600, 400


def _raw_image() -> np.ndarray:
    """Frame with a bright block exactly where the patient sticker lives, so a
    test can tell the masked region from the unmasked one."""
    img = np.zeros((_H, _W, 3), np.uint8)
    x, y, w, h = geometry_for(MAXX_ORTHO).patient_region.to_pixels(_W, _H)
    img[y:y + h, x:x + w] = 255
    return img


def _jpeg_bytes() -> bytes:
    ok, buf = cv2.imencode(".jpg", _raw_image())
    assert ok
    return buf.tobytes()


def _fake_client(initials="JD", confidence="high"):
    client = MagicMock()
    client.messages.create.return_value = SimpleNamespace(
        content=[SimpleNamespace(
            type="text",
            text='{"initials": "%s", "confidence": "%s"}' % (initials, confidence),
        )],
    )
    return client


def _flag_on(fake_client):
    """Flag on + a stubbed Anthropic client."""
    return (
        patch.object(patient.settings, "extract_patient_initials", True),
        patch.object(patient.settings, "anthropic_api_key", "sk-test"),
        patch.object(patient.settings, "offline_mode", False),
        patch("anthropic.Anthropic", return_value=fake_client),
    )


# ---------------------------------------------------------------------------
# Default posture: the flag is off, so no patient data is read at all.
# ---------------------------------------------------------------------------
def test_flag_is_off_by_default():
    from app.config import Settings
    assert Settings().extract_patient_initials is False


def test_flag_off_reads_nothing_and_calls_no_api():
    with patch("anthropic.Anthropic") as ctor:
        result = patient.extract_initials(_raw_image(), MAXX_ORTHO)
    assert result == {"value": None, "confidence": "low"}
    ctor.assert_not_called()


# ---------------------------------------------------------------------------
# The redaction gate still comes first. A ticket that fails it sends nothing.
# ---------------------------------------------------------------------------
def test_gate_failure_never_reads_initials():
    batch = db.create_batch()
    with patch.object(run, "redact_patient_region", return_value=(None, False)), \
         patch.object(run.patient, "extract_initials") as spy:
        res = ingest_image(_jpeg_bytes(), "MO-gate.jpg", batch["id"])

    assert res["status"] == "manual_queue"
    spy.assert_not_called()


def test_encode_failure_never_reads_initials():
    batch = db.create_batch()
    redacted = np.zeros((_H, _W, 3), np.uint8)
    with patch.object(run, "redact_patient_region", return_value=(redacted, True)), \
         patch.object(run.preprocess, "encode_image", return_value=b""), \
         patch.object(run.patient, "extract_initials") as spy:
        res = ingest_image(_jpeg_bytes(), "MO-encode.jpg", batch["id"])

    assert res["status"] == "manual_queue"
    spy.assert_not_called()


# ---------------------------------------------------------------------------
# Minimum necessary: only the sticker crop leaves the process.
# ---------------------------------------------------------------------------
def test_only_the_patient_crop_is_sent_never_the_whole_ticket():
    client = _fake_client()
    a, b, c, d = _flag_on(client)
    with a, b, c, d:
        patient.extract_initials(_raw_image(), MAXX_ORTHO)

    content = client.messages.create.call_args.kwargs["messages"][0]["content"]
    image_block = next(b for b in content if b["type"] == "image")

    import base64
    sent = cv2.imdecode(
        np.frombuffer(base64.standard_b64decode(image_block["source"]["data"]), np.uint8),
        cv2.IMREAD_COLOR,
    )
    x, y, w, h = geometry_for(MAXX_ORTHO).patient_region.to_pixels(_W, _H)
    assert sent.shape[:2] == (h, w), "must send the sticker crop, not the full frame"
    assert sent.shape[:2] != (_H, _W)


def test_initials_call_runs_at_low_effort():
    """Two letters off a crop needs no deliberation. Without this, Sonnet 5
    would default to "high" and we'd pay for reasoning we don't need."""
    client = _fake_client()
    a, b, c, d = _flag_on(client)
    with a, b, c, d:
        patient.extract_initials(_raw_image(), MAXX_ORTHO)

    assert client.messages.create.call_args.kwargs["output_config"] == {
        "effort": "low"}


def test_stored_image_is_still_fully_redacted():
    batch = db.create_batch()
    stored: dict = {}

    def capture(bucket, name, data, content_type):
        stored["bytes"] = data
        return f"{bucket}/{name}"

    client = _fake_client()
    a, b, c, d = _flag_on(client)
    with a, b, c, d, patch.object(run, "put_object", side_effect=capture):
        res = ingest_image(_jpeg_bytes(), "MO-stored.jpg", batch["id"])

    assert res["status"] == "pending_review"
    img = cv2.imdecode(np.frombuffer(stored["bytes"], np.uint8), cv2.IMREAD_COLOR)
    x, y, w, h = geometry_for(MAXX_ORTHO).patient_region.to_pixels(_W, _H)
    region = img[y:y + h, x:x + w]
    # The raw frame had this block at 255; the stored one must be masked.
    assert region.mean() < 100, "patient sticker must still be masked in storage"


# ---------------------------------------------------------------------------
# Only ever two letters come back — the schema is the safety net.
# ---------------------------------------------------------------------------
def test_clean_accepts_only_two_letters():
    assert patient._clean("JD") == "JD"
    assert patient._clean("j d") == "JD"
    assert patient._clean("J.D.") == "JD"
    assert patient._clean("John Doe") is None      # a full name is rejected
    assert patient._clean("JMD") is None           # middle initial slipped in
    assert patient._clean("J") is None
    assert patient._clean(None) is None


def test_full_name_in_the_response_is_discarded():
    client = _fake_client(initials="John Doe")
    a, b, c, d = _flag_on(client)
    with a, b, c, d:
        result = patient.extract_initials(_raw_image(), MAXX_ORTHO)
    assert result == {"value": None, "confidence": "low"}


def test_api_error_leaves_the_cell_blank():
    client = MagicMock()
    client.messages.create.side_effect = Exception("api down")
    a, b, c, d = _flag_on(client)
    with a, b, c, d:
        result = patient.extract_initials(_raw_image(), MAXX_ORTHO)
    assert result == {"value": None, "confidence": "low"}


def test_unknown_template_reads_nothing():
    client = _fake_client()
    a, b, c, d = _flag_on(client)
    with a, b, c, d:
        result = patient.extract_initials(_raw_image(), "Unknown")
    assert result == {"value": None, "confidence": "low"}
    client.messages.create.assert_not_called()


# ---------------------------------------------------------------------------
# End to end: the initials reach the Inits column of the Usage sheet.
# ---------------------------------------------------------------------------
def test_initials_land_in_the_inits_column():
    batch = db.create_batch()
    client = _fake_client(initials="AB")
    a, b, c, d = _flag_on(client)
    with a, b, c, d:
        res = ingest_image(_jpeg_bytes(), "MO-inits.jpg", batch["id"])

    assert res["status"] == "pending_review"
    ticket = db.get_ticket(res["ticket_id"])
    assert ticket["patient_initials"] == "AB"
    assert ticket["patient_initials_conf"] == "high"

    def _f(value, confidence="high"):
        return {"value": value, "confidence": confidence}

    assemble_and_persist(ticket, {
        "header": {"surgeon": _f("Woodworth"), "rep_code": _f("GR-ME-001"),
                   "surgery_date": _f("2026-06-01")},
        "lines": [{"index": 0, "ref": _f("INIT-REF-1"), "qty": _f(1),
                   "unit_price": _f(100)}],
        "freight": _f(None, "low"), "grand_total": _f(100),
    }, [{"gtin": None, "lot": None, "expiry": None, "mfg": None, "serial": None,
         "raw": None, "decoded": False, "ref": None}])

    ws = load_workbook(io.BytesIO(write_review_workbook(batch["id"])))["Usage"]
    headers = [c.value for c in ws[1]]
    assert headers[3] == "Inits", "Inits sits right after Surgeon"
    row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(headers)}
    assert row["Inits"] == "AB"
    assert row["Surgeon"] == "Woodworth"


def test_corrected_workbook_inits_column_round_trips():
    """A corrected workbook carrying the Inits column parses it back out. The
    parser matches on header name, so the column's position doesn't matter."""
    from app.sheets.read import parse_corrected_workbook

    batch = db.create_batch()
    ticket = db.create_ticket({
        "batch_id": batch["id"], "source_filename": "MO-rt.jpg",
        "status": "pending_review", "patient_initials": "EF",
        "patient_initials_conf": "high",
    })

    def _f(value, confidence="high"):
        return {"value": value, "confidence": confidence}

    assemble_and_persist(db.get_ticket(ticket["ticket_id"]), {
        "header": {"surgeon": _f("Woodworth"), "surgery_date": _f("2026-06-01")},
        "lines": [], "freight": _f(None, "low"), "grand_total": _f(None, "low"),
    }, [])

    parsed = parse_corrected_workbook(write_review_workbook(batch["id"]))
    assert parsed["tickets"][ticket["ticket_id"]]["patient_initials"] == "EF"


def test_reprocessing_keeps_the_initials():
    """assemble runs on the redacted photo and must not blank a stored value."""
    batch = db.create_batch()
    ticket = db.create_ticket({
        "batch_id": batch["id"], "source_filename": "MO-reproc.jpg",
        "status": "pending_review", "patient_initials": "CD",
        "patient_initials_conf": "medium",
    })

    def _f(value, confidence="high"):
        return {"value": value, "confidence": confidence}

    for _ in range(2):
        assemble_and_persist(db.get_ticket(ticket["ticket_id"]), {
            "header": {"surgeon": _f("Nobody"), "surgery_date": _f("2026-06-01")},
            "lines": [], "freight": _f(None, "low"), "grand_total": _f(None, "low"),
        }, [])

    assert db.get_ticket(ticket["ticket_id"])["patient_initials"] == "CD"
