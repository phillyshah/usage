"""Step 5 — filling blank Price cells from the hospital price list.

Covers the twelve named acceptance tests from the work instructions, plus the
regressions that the real price-list workbook and the real part master turned
up while building this:

  * the ``***`` fill is NOT three characters (``UFCR***-GK`` is ``UFCRLA00-GK``),
  * the component tiers must short-circuit or ``MTUUX`` poisons ``MTUUX***-GK``,
  * the two tab layouts have different numbers of metadata columns,
  * Excel writes ``800002`` as ``800002.0``,
  * a formula that displays empty is not a blank cell,
  * a wasted (yellow) Price cell can also be blank, and must stay yellow.
"""
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.db import db
from app.main import app
from app.pipeline.assemble import assemble_and_persist
from app.pricing import match as mt
from app.pricing import normalize as nz
from app.pricing.enrich import EnrichmentError, enrich_workbook
from app.pricing.ingest import parse_price_list
from app.sheets.write import write_review_workbook

client = TestClient(app)

MH_TAB = "MH for MO"
MO_TAB = "Summary Price List"


# --------------------------------------------------------------------------
# Fixtures / builders
# --------------------------------------------------------------------------
def _f(value, confidence="high"):
    return {"value": value, "confidence": confidence}


def _empty_label():
    return {"gtin": None, "lot": None, "expiry": None, "mfg": None,
            "serial": None, "raw": None, "decoded": False, "ref": None}


def price_list_bytes(tabs: dict) -> bytes:
    """tabs -> {tab: {"meta": [...], "hospitals": [...],
                      "rows": [(item, [meta values...], {hospital: price})]}}

    Row 1 is junk and row 2 is the header, mirroring the real workbook.
    """
    wb = Workbook()
    wb.remove(wb.active)
    for name, spec in tabs.items():
        ws = wb.create_sheet(name)
        ws.append([spec.get("banner", "PRICE LIST")])
        ws.append(list(spec["meta"]) + list(spec["hospitals"]))
        for item, meta, prices in spec["rows"]:
            ws.append([item] + list(meta)
                      + [prices.get(h) for h in spec["hospitals"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def seed_price_list(tabs: dict) -> None:
    db.replace_hospital_prices(parse_price_list(price_list_bytes(tabs))["rows"])


def seed_usage(rows: list[dict]) -> bytes:
    """rows -> the real generated workbook, so the column contract can't drift.

    Each row: {filename, entity, hospital, ref, price, wasted, description,
    part_type, category}.

    The surgeon and part masters are seeded from the rows because assemble drops
    what they don't cover: an unknown REF leaves Ref Number blank, and an
    unmatched surgeon leaves Hospital blank — the same two upstream holes that
    make most real rows unpriceable.
    """
    db.replace_reference_surgeons([
        {"surgeon_distcode": f"PRICERPR-{i}", "surgeon_last_name": "Pricer",
         "dist_code": f"PR-{i}", "status": "Active", "surgeon_full_name": "P Pricer",
         "hospital": r["hospital"], "region": "X", "distributor_rep": "R"}
        for i, r in enumerate(rows) if r.get("hospital")
    ])
    db.replace_reference_part_info([
        {"part_number": r["ref"],
         "description": r.get("description", "TIBIAL BASE PLATE (TITAN)"),
         "part_type": r.get("part_type", "Tibial"),
         "category": r.get("category", "Knee")}
        for r in rows if r.get("ref")
    ])
    batch = db.create_batch()
    for i, r in enumerate(rows):
        ticket = db.create_ticket({
            "batch_id": batch["id"], "entity": r.get("entity"),
            "source_filename": r["filename"], "surgeon": "Pricer",
            "rep_code": f"PR-{i}", "hospital": r.get("hospital"),
            "surgery_date": "2026-06-01", "status": "pending_review",
        })
        price = r.get("price")
        vision = {
            "header": {
                "surgeon": _f("Pricer"), "rep_code": _f(f"PR-{i}"),
                "hospital": _f(r["hospital"]) if r.get("hospital") else _f(None, "low"),
                "surgery_date": _f("2026-06-01"),
                "entity": _f(r["entity"]) if r.get("entity") else _f(None, "low"),
            },
            "lines": [{
                "index": 0,
                "ref": _f(r["ref"]) if r.get("ref") else _f(None, "low"),
                "lot": _f(f"L{i}"), "qty": _f(1),
                "unit_price": _f(price) if price is not None else _f(None, "low"),
                "wasted": _f(bool(r.get("wasted"))),
            }],
            "freight": _f(None, "low"), "grand_total": _f(None, "low"),
        }
        assemble_and_persist(ticket, vision, [_empty_label()])
    return write_review_workbook(batch["id"])


def usage_prices(data: bytes) -> list[tuple]:
    """[(row, value, fill_rgb)] for every Usage data row, in sheet order."""
    ws = load_workbook(io.BytesIO(data))["Usage"]
    headers = [c.value for c in ws[1]]
    col = headers.index("Price") + 1
    out = []
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=col)
        rgb = getattr(cell.fill.fgColor, "rgb", None)
        out.append((r, cell.value, rgb))
    return out


@pytest.fixture(autouse=True)
def _clean_price_list():
    db.replace_hospital_prices([])
    db.replace_reference_part_info([])
    db.replace_reference_surgeons([])
    yield


# ==========================================================================
# Acceptance tests 1-4 — hospital normalization and matching
# ==========================================================================
def test_at1_ctr_expands_to_center_midname():
    """AT1: 'Seaside Surg Ctr Group' must equal 'Seaside Surgery Center Group'."""
    assert (nz.normalize_hospital("Seaside Surg Ctr Group")
            == nz.normalize_hospital("Seaside Surgery Center Group"))
    m = mt.match_hospital("Seaside Surg Ctr Group", ["Seaside Surgery Center Group"])
    assert m.method == "exact" and m.confident


def test_at2_legal_suffix_and_extra_spaces_normalize():
    """AT2: 'Acme Hospital, LLC' == 'Acme  Hospital' (suffix + whitespace)."""
    assert (nz.normalize_hospital("Acme Hospital, LLC")
            == nz.normalize_hospital("Acme  Hospital"))


def test_at3_fuzzy_single_candidate_matches():
    """AT3: one credible candidate at >=75% resolves."""
    m = mt.match_hospital("AdventHealth Carrolwood", ["AdventHealth Carrollwood"])
    assert m.matched and m.name == "AdventHealth Carrollwood"


def test_at4_multiple_credible_candidates_are_ambiguous():
    """AT4: two candidates that score alike resolve to nothing, not a coin flip."""
    m = mt.match_hospital("Mercy Surgery Institute",
                          ["Mercy Surgical Institute A", "Mercy Surgical Institute B"])
    assert m.method == "ambiguous" and not m.matched


def test_near_duplicate_is_not_ambiguous_when_one_clearly_wins():
    """The Summary tab lists 'Advanced Surg Ctr of North County' AND
    '... North County HIgh Demand'. Both clear 75%; calling that a tie would
    throw away a match that is effectively exact."""
    m = mt.match_hospital(
        "Advanced Surgery Center of North County",
        ["Advanced Surg Ctr of North County",
         "Advanced Surg Ctr of North County HIgh Demand"])
    assert m.name == "Advanced Surg Ctr of North County"


def test_fuzzy_hospital_never_backs_a_green_price():
    """A fuzzy hospital match prices the row, but as an estimate. Green says
    'came straight from the list, don't check it' and fuzzy can't promise that."""
    m = mt.match_hospital("AdventHealth Carrolwood", ["AdventHealth Carrollwood"])
    assert m.matched and m.method == "fuzzy" and not m.confident


def test_health_system_parenthetical_is_stripped():
    assert (nz.normalize_hospital("Centerpoint Med Ctr (HCA)")
            == nz.normalize_hospital("Centerpoint Medical Center"))


# ==========================================================================
# Acceptance test 9 + component matching
# ==========================================================================
def test_at9_cr_and_ps_never_match_on_description():
    """AT9: a CR component and a PS component are different devices however
    similar the rest of the text reads."""
    catalog = {"MLPSX": "Tibial Articular Surface PS"}
    m = mt.match_component(None, "Tibial Articular Surface CR", catalog)
    assert not m.matched


def test_wildcard_fill_is_not_fixed_at_three_characters():
    """``UFCR***-GK`` stands for the real REF ``UFCRLA00-GK`` — a four-character
    fill. A ``.{3}`` quantifier silently drops the whole family."""
    pat = mt.compile_code_pattern("UFCR***-GK")
    assert pat.match("UFCRLA00-GK")
    assert mt.compile_code_pattern("ACLMR***-UK").match("ACLMRL100-UK")


def test_wildcard_does_not_cross_k_and_gk_suffixes():
    assert not mt.compile_code_pattern("MTUUX***-GK").match("MTUUX100-K")


def test_lowercase_xxx_is_a_wildcard_and_uppercase_xx_is_literal():
    """``311xxx`` is a code family; ``DAXX00D-F`` is a real part number."""
    assert mt.compile_code_pattern("311xxx").match("311541")
    assert not mt.compile_code_pattern("311xxx").match("314541")
    assert mt.compile_code_pattern("DAXX00D-F") is None


def test_tier_shortcircuit_prevents_the_mtuux_false_conflict():
    """The bare code ``MTUUX`` is a prefix of ``MTUUX100-GK``, which the pattern
    ``MTUUX***-GK`` also matches — at different prices. Unioning the tiers would
    read that as a conflict and refuse to price the row."""
    catalog = {"MTUUX***-GK": "TIBIAL BASE PLATE (TITAN)", "MTUUX": "TBP"}
    m = mt.match_component("MTUUX100-GK", None, catalog)
    assert m.tier == "wildcard" and m.codes == ("MTUUX***-GK",)
    # and the bare code still wins for a REF the pattern doesn't cover
    assert mt.match_component("MTUUX100-K", None, catalog).codes == ("MTUUX",)


def test_prefix_tier_resolves_a_family_written_without_wildcards():
    """'MO-MSFC' and 'ALCRX' are written as bare codes but are prefixes of the
    real REFs ('MO-MSFC-46/MB', 'ALCRXA109-K')."""
    catalog = {"MO-MSFC": "Femoral", "ALCRX": "All Poly Tibial CR"}
    assert mt.match_component("MO-MSFC-46/MB", None, catalog).tier == "prefix"
    assert mt.match_component("ALCRXA109-K", None, catalog).codes == ("ALCRX",)


def test_generic_description_words_alone_cannot_carry_a_match():
    catalog = {"ZZZ": "Femoral Component"}
    assert not mt.match_component(None, "Tibial Component", catalog).matched


# ==========================================================================
# Ingest — the two real tab layouts and Excel's float coercion
# ==========================================================================
def test_parses_both_meta_column_layouts():
    """'MH for MO' is Item/Description (2 meta columns); 'Summary Price List'
    is Item/Class/Part Type (3). The first hospital column is detected."""
    parsed = parse_price_list(price_list_bytes({
        MH_TAB: {"meta": ["Item", "Description"], "hospitals": ["Blake Hospital"],
                 "rows": [("ALCRX", ["All Poly Tibial CR"], {"Blake Hospital": 1000})]},
        MO_TAB: {"meta": ["Item", "Class", "Part Type"],
                 "hospitals": ["River Surgical Institute"],
                 "rows": [("MTUUX***-GK", ["Knee", "Tibial"],
                           {"River Surgical Institute": 925})]},
    }))
    rows = {(r["tab"], r["item_code"]): r for r in parsed["rows"]}
    assert rows[(MH_TAB, "ALCRX")]["hospital"] == "Blake Hospital"
    assert rows[(MH_TAB, "ALCRX")]["unit_price"] == 1000
    assert rows[(MO_TAB, "MTUUX***-GK")]["unit_price"] == 925
    assert parsed["tabs"][MO_TAB]["hospitals"] == 1


def test_excel_float_item_code_normalizes_to_the_real_ref():
    """Excel stores the numeric part number 800002 as 800002.0."""
    parsed = parse_price_list(price_list_bytes({
        MH_TAB: {"meta": ["Item", "Description"], "hospitals": ["Blake Hospital"],
                 "rows": [(800002.0, ["Screw"], {"Blake Hospital": 75})]},
    }))
    assert parsed["rows"][0]["item_code"] == "800002"


def test_duplicate_hospital_columns_agreeing_are_kept():
    """'Surgcenter of Plano' appears twice on the real Summary tab. A
    header-keyed parser would silently drop one."""
    parsed = parse_price_list(price_list_bytes({
        MO_TAB: {"meta": ["Item", "Class", "Part Type"],
                 "hospitals": ["Surgcenter of Plano", "Surgcenter of Plano"],
                 "rows": [("ALCRX", ["Knee", "Tibial"],
                           {"Surgcenter of Plano": 400})]},
    }))
    assert len(parsed["rows"]) == 1
    assert parsed["rows"][0]["unit_price"] == 400


def test_a_workbook_with_no_prices_is_rejected():
    with pytest.raises(ValueError):
        parse_price_list(price_list_bytes({
            MH_TAB: {"meta": ["Item", "Description"], "hospitals": ["Blake Hospital"],
                     "rows": [("ALCRX", ["x"], {})]},
        }))


# ==========================================================================
# Acceptance tests 5-8, 10-12 — the end-to-end run
# ==========================================================================
def _mh_list(prices=None):
    return {MH_TAB: {"meta": ["Item", "Description"],
                     "hospitals": ["Blake Hospital (HCA)"],
                     "rows": [("MTUUX***-GK", ["TIBIAL BASE PLATE (TITAN)"],
                               prices or {"Blake Hospital (HCA)": 925})]}}


def test_at5_direct_match_fills_neon_green():
    seed_price_list(_mh_list())
    data = seed_usage([{"filename": "MH17469.jpg", "entity": "Maxx Health",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    out, summary = enrich_workbook(data)
    assert summary["direct"] == 1 and summary["estimates"] == 0
    _, value, rgb = usage_prices(out)[0]
    assert value == 925
    assert rgb.endswith("39FF14")


def test_at6_estimate_fills_rose():
    """No price for this hospital, but the same REF is priced elsewhere in the
    same workbook at the same hospital — rung 1 of the estimate ladder."""
    seed_price_list(_mh_list({"Blake Hospital (HCA)": 925}))
    data = seed_usage([
        {"filename": "MH1.jpg", "entity": "Maxx Health",
         "hospital": "Nowhere Surgical Partners", "ref": "ZZ-NOT-LISTED", "price": 640},
        {"filename": "MH2.jpg", "entity": "Maxx Health",
         "hospital": "Nowhere Surgical Partners", "ref": "ZZ-NOT-LISTED"},
    ])
    out, summary = enrich_workbook(data)
    assert summary["estimates"] == 1 and summary["direct"] == 0
    filled = [p for p in usage_prices(out) if p[1] == 640 and p[2]]
    assert any(rgb.endswith("FFC7CE") for _, _, rgb in filled)


def test_at7_an_existing_zero_price_is_left_alone():
    """A zero is a real price of zero, not a hole."""
    seed_price_list(_mh_list())
    data = seed_usage([{"filename": "MH1.jpg", "entity": "Maxx Health",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK",
                        "price": 0}])
    out, summary = enrich_workbook(data)
    assert summary["eligible"] == 0 and summary["direct"] == 0
    assert usage_prices(out)[0][1] == 0


def test_at8_a_wasted_yellow_cell_is_skipped_and_stays_yellow():
    """write.py paints a wasted line's Price yellow even when the price is blank.
    A wasted component's price is a business decision, not a lookup."""
    seed_price_list(_mh_list())
    data = seed_usage([{"filename": "MH1.jpg", "entity": "Maxx Health",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK",
                        "wasted": True}])
    out, summary = enrich_workbook(data)
    assert summary["skipped_wasted"] == 1
    assert summary["eligible"] == 0
    _, value, rgb = usage_prices(out)[0]
    assert value is None and rgb.endswith("FFFF00")


def test_at10_and_at11_a_row_prefers_its_own_distributors_tab():
    """AT10/AT11, as amended.

    The written instructions said a Maxx Health row may use *only* the MH tab.
    The first real workbook showed why that is too strong: 27 of its 88 blank
    prices belonged to hospitals listed only on a tab the rule forbade, so the
    tabs are where an account is listed, not a price agreement scoped to one
    distributor. The rule is now preference, not exclusivity — and the part that
    still matters is asserted here: where BOTH tabs know the hospital, the
    ticket's own tab wins and its price is the one used."""
    seed_price_list({
        MH_TAB: {"meta": ["Item", "Description"], "hospitals": ["Blake Hospital (HCA)"],
                 "rows": [("MTUUX***-GK", ["TIBIAL BASE PLATE (TITAN)"],
                           {"Blake Hospital (HCA)": 925})]},
        MO_TAB: {"meta": ["Item", "Class", "Part Type"],
                 "hospitals": ["Blake Hospital (HCA)"],
                 "rows": [("MTUUX***-GK", ["Knee", "Tibial"],
                           {"Blake Hospital (HCA)": 5555})]},
    })
    data = seed_usage([{"filename": "MH17469.jpg", "entity": "Maxx Health",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    out, summary = enrich_workbook(data)
    assert summary["tabs"] == [MH_TAB]
    assert usage_prices(out)[0][1] == 925


def test_at12_a_failed_run_leaves_the_previous_output_in_place():
    seed_price_list(_mh_list())
    good = seed_usage([{"filename": "MH17469.jpg", "entity": "Maxx Health",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    ok = client.post("/pricing/enrich", files={"file": ("usage.xlsx", io.BytesIO(good),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert ok.status_code == 200
    first_run = ok.json()["run_id"]

    bad = client.post("/pricing/enrich", files={"file": ("junk.xlsx", io.BytesIO(b"not a workbook"),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert bad.status_code == 400
    assert bad.json()["last_good"]["run_id"] == first_run
    assert client.get(f"/pricing/runs/{first_run}/sheet").status_code == 200
    assert client.get("/pricing/latest").json()["run_id"] == first_run


# ==========================================================================
# Workbook-integrity guarantees
# ==========================================================================
def test_an_entity_the_tab_table_does_not_cover_fails_the_run():
    seed_price_list(_mh_list())
    data = seed_usage([{"filename": "ticket-001.jpg", "entity": "Some Other Distributor",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    with pytest.raises(EnrichmentError) as exc:
        enrich_workbook(data)
    assert exc.value.reason == "unconfigured_distributor"


def test_a_blank_entity_falls_back_to_the_mh_mo_filename_prefix():
    """write.py blanks Entity whenever the vision read was low-confidence, so
    the filename convention has to be a real fallback, not a nicety."""
    seed_price_list(_mh_list())
    data = seed_usage([{"filename": "MH17469.jpg", "entity": None,
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    _, summary = enrich_workbook(data)
    assert summary["tabs"] == [MH_TAB] and summary["direct"] == 1


def test_formulas_are_never_treated_as_blank_and_survive_the_round_trip():
    seed_price_list(_mh_list())
    data = seed_usage([{"filename": "MH1.jpg", "entity": "Maxx Health",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Usage"]
    col = [c.value for c in ws[1]].index("Price") + 1
    ws.cell(row=2, column=col).value = '=IF(1=1,"","")'
    buf = io.BytesIO()
    wb.save(buf)

    out, summary = enrich_workbook(buf.getvalue())
    assert summary["eligible"] == 0
    assert usage_prices(out)[0][1] == '=IF(1=1,"","")'


def test_a_workbook_with_charts_or_images_is_refused_rather_than_stripped():
    """openpyxl cannot round-trip them, and silently destroying the operator's
    work is worse than refusing to touch the file."""
    wb = Workbook()
    wb.active.title = "Usage"
    wb.active.append(["Source Image Filename", "Hospital", "Price", "Ref Number"])
    buf = io.BytesIO()
    wb.save(buf)
    import zipfile
    doctored = io.BytesIO()
    with zipfile.ZipFile(buf, "r") as src, zipfile.ZipFile(doctored, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/charts/chart1.xml", "<c:chartSpace/>")
    with pytest.raises(EnrichmentError) as exc:
        enrich_workbook(doctored.getvalue())
    assert exc.value.reason == "unsupported_workbook_content"


def test_a_workbook_without_a_usage_sheet_is_rejected():
    wb = Workbook()
    wb.active.title = "Something Else"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(EnrichmentError) as exc:
        enrich_workbook(buf.getvalue())
    assert exc.value.reason == "missing_usage_columns"


def test_a_missing_tab_fails_the_run_rather_than_pricing_from_another():
    db.replace_hospital_prices([])
    data = seed_usage([{"filename": "MH17469.jpg", "entity": "Maxx Health",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    with pytest.raises(EnrichmentError) as exc:
        enrich_workbook(data)
    assert exc.value.reason == "missing_tab"


def test_unresolved_rows_keep_their_red_fill_and_are_reported_by_cause():
    """Most blanks on a real run are blank because the Hospital or Ref cell
    upstream is itself blank — without the breakdown this looks like a bug."""
    seed_price_list(_mh_list())
    data = seed_usage([{"filename": "MH1.jpg", "entity": "Maxx Health",
                        "hospital": None, "ref": "MTUUX100-GK"}])
    out, summary = enrich_workbook(data)
    assert summary["unresolved"] == 1
    assert summary["unresolved_causes"] == {"no_hospital": 1}
    _, value, rgb = usage_prices(out)[0]
    assert value is None and rgb.endswith("F4CCCC")


def test_the_run_summary_always_adds_up():
    seed_price_list(_mh_list())
    data = seed_usage([
        {"filename": "MH1.jpg", "entity": "Maxx Health",
         "hospital": "Blake Hospital", "ref": "MTUUX100-GK"},
        {"filename": "MH2.jpg", "entity": "Maxx Health",
         "hospital": "Blake Hospital", "ref": "ZZ-UNKNOWN"},
        {"filename": "MH3.jpg", "entity": "Maxx Health",
         "hospital": "Blake Hospital", "ref": "MTUUX200-GK", "price": 400},
    ])
    _, s = enrich_workbook(data)
    assert s["eligible"] == s["direct"] + s["estimates"] + s["unresolved"]


def test_the_price_list_appears_in_reference_status():
    seed_price_list(_mh_list())
    masters = client.get("/reference/status").json()["masters"]
    assert masters["prices"]["rows"] >= 1


# ==========================================================================
# Regressions found by running the real Hospital_Price_List against the code.
# Every one of these failed silently: no exception, no wrong number on screen,
# just a price that never appeared. That is why a green suite missed them.
# ==========================================================================
def test_four_star_wildcards_match_like_three_star_ones():
    """The catalogue writes three, four or more stars — ``ACLM****-UK`` and
    ``RFPS****-GK`` are both real. Splitting on a literal ``***`` leaves the
    fourth star to be escaped into the pattern, which then matches nothing;
    19 price rows on the Summary tab were unreachable."""
    assert mt.compile_code_pattern("ACLM****-UK").match("ACLMRL100-UK")
    assert mt.compile_code_pattern("RFPS****-GK").match("RFPSLA00-GK")
    # and the three-star families still behave
    assert mt.compile_code_pattern("MTUUX***-GK").match("MTUUX100-GK")
    assert not mt.compile_code_pattern("MTUUX***-GK").match("MTUUX100-K")


def test_aggregate_columns_are_not_ingested_as_hospitals():
    """The Summary tab's last column is ``AVERAGE ITEM PRICE`` — a spreadsheet
    statistic. Left in, the estimate ladder medians *across hospitals* and folds
    that average back in as though it were an independent account."""
    parsed = parse_price_list(price_list_bytes({
        MO_TAB: {"meta": ["Item", "Class", "Part Type"],
                 "hospitals": ["Blake Hospital", "AVERAGE ITEM PRICE"],
                 "rows": [("ALCRX", ["Knee", "Tibial"],
                           {"Blake Hospital": 1000, "AVERAGE ITEM PRICE": 1630})]},
    }))
    assert [r["hospital"] for r in parsed["rows"]] == ["Blake Hospital"]
    assert parsed["tabs"][MO_TAB]["hospitals"] == 1


def test_an_average_column_cannot_be_matched_as_a_hospital():
    assert not mt.match_hospital("Average Item Price", ["Blake Hospital"]).matched


def test_a_generic_system_name_does_not_pick_one_sibling_facility():
    """'Baylor, Scott, & White' names no facility, and the list holds five of
    them. Picking the highest scorer means picking on suffix length."""
    m = mt.match_hospital("Baylor, Scott, & White", [
        "Baylor Scott & White Star", "Baylor Scott & White Frisco",
        "Baylor Scott & White Sherman", "Baylor Scott & White Sunnyvale",
        "Baylor Scott & White Centennial"])
    assert m.method == "ambiguous" and not m.matched


def test_a_named_facility_resolves_against_its_siblings():
    """The same five siblings, but this query *does* name one. Stripping the
    generic words makes the two forms identical."""
    m = mt.match_hospital("Baylor Scott & White Medical Center - Sunnyvale", [
        "Baylor Scott & White Sunnyvale", "Baylor Scott & White Centennial"])
    assert m.method == "core"
    assert m.name == "Baylor Scott & White Sunnyvale"


def test_a_typo_resolves_when_it_is_near_exact():
    """Hazelton / Hazleton, against two other Lehigh Valley facilities."""
    m = mt.match_hospital("Lehigh Valley Hospital - Hazelton", [
        "Lehigh Valley Hosp Hazleton", "Lehigh Valley Hosp Highland",
        "Lehigh Valley Hosp Pocono"])
    assert m.name == "Lehigh Valley Hosp Hazleton"


def test_core_matches_are_not_green():
    """``meaningful_hospital`` strips 'hospital', 'surgery' and 'center', so a
    hospital and its surgery centre collapse together — plausibly two different
    accounts. Good enough to estimate from, not good enough to tell the reviewer
    not to check."""
    m = mt.match_hospital("Boca Raton Hospital", ["Boca Raton Surg Ctr"])
    assert m.method == "core" and m.matched
    assert not m.confident


def test_family_price_at_the_same_hospital_is_used_when_the_variant_is_unpriced():
    """The component tier short-circuits, so a REF ending -GK resolves to the
    wildcard row. When that row has no price at this hospital but the bare
    family row does, the family price is the best same-hospital evidence there
    is — written rose, because the two rows are demonstrably different prices."""
    seed_price_list({MH_TAB: {
        "meta": ["Item", "Description"],
        "hospitals": ["Blake Hospital (HCA)", "Other Hospital"],
        "rows": [
            # the variant row is priced somewhere else, but not here
            ("MTUUX***-GK", ["TIBIAL BASE PLATE (TITAN)"], {"Other Hospital": 925}),
            # the family row is priced HERE
            ("MTUUX", ["TIBIAL BASE PLATE"], {"Blake Hospital (HCA)": 700}),
        ]}})
    data = seed_usage([{"filename": "MH1.jpg", "entity": "Maxx Health",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    out, summary = enrich_workbook(data)
    assert summary["direct"] == 0 and summary["estimates"] == 1
    _, value, rgb = usage_prices(out)[0]
    assert value == 700
    assert rgb.endswith("FFC7CE")


def test_the_variant_row_still_wins_when_both_are_priced_here():
    """The fall-through must not become a general union: where both rows carry a
    price for this hospital, the more specific one is the answer, and it is a
    direct (green) price."""
    seed_price_list({MH_TAB: {
        "meta": ["Item", "Description"],
        "hospitals": ["Blake Hospital (HCA)"],
        "rows": [
            ("MTUUX***-GK", ["TIBIAL BASE PLATE (TITAN)"],
             {"Blake Hospital (HCA)": 925}),
            ("MTUUX", ["TIBIAL BASE PLATE"], {"Blake Hospital (HCA)": 700}),
        ]}})
    data = seed_usage([{"filename": "MH1.jpg", "entity": "Maxx Health",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    out, summary = enrich_workbook(data)
    assert summary["direct"] == 1
    _, value, rgb = usage_prices(out)[0]
    assert value == 925
    assert rgb.endswith("39FF14")


# ==========================================================================
# Regressions from the first real usage workbook (review_542abe77.xlsx):
# 153 rows, 88 blank prices, and it failed outright on row 1.
# ==========================================================================
@pytest.mark.parametrize("entity", [
    "Maxx Orthopedics",        # 30 tickets in the real file
    "Maxx Orthopedics, Inc",   # 6
    "Maxx Orthopedics, Inc.",  # 1
])
def test_entity_spelling_variants_all_resolve(entity):
    """Entity is the model's reading of printed text, so one batch from one
    company arrives spelled several ways. An exact-string lookup fails on all
    but the first, and one failing row used to abort the whole workbook."""
    from app.pricing.enrich import resolve_entity
    _, tab = resolve_entity("ticket-1", {"ticket-1": entity})
    assert tab == MO_TAB


def test_the_filename_resolves_the_entity_when_the_text_is_unmappable():
    """One real ticket's Entity read just 'MAXX' — ambiguous between the two
    companies, and no amount of normalising fixes it. Its filename said
    MO18711-A, which is not ambiguous at all."""
    from app.pricing.enrich import resolve_entity
    assert resolve_entity("MO18711-A", {"MO18711-A": "MAXX"})[1] == MO_TAB
    # and with the Entity column blank, which write.py does on low confidence
    assert resolve_entity("MO18711-A", {})[1] == MO_TAB


def test_the_filename_wins_when_it_disagrees_with_the_ticket_text():
    """detect_template is already trusted to pick the PHI redaction region off
    the filename; pricing should not trust a different signal more."""
    from app.pricing.enrich import resolve_entity
    entity, tab = resolve_entity("MO18711-A", {"MO18711-A": "Maxx Health"})
    assert (entity, tab) == ("Maxx Orthopedics", MO_TAB)


def test_one_unresolvable_row_does_not_fail_the_run():
    """A 153-row workbook used to die on a single ticket whose distributor could
    not be identified."""
    seed_price_list(_mh_list())
    data = seed_usage([
        {"filename": "MH1.jpg", "entity": "Maxx Health",
         "hospital": "Blake Hospital", "ref": "MTUUX100-GK"},
        {"filename": "ticket-oddly-named", "entity": "Someone Else Entirely",
         "hospital": "Blake Hospital", "ref": "MTUUX100-GK"},
    ])
    _, summary = enrich_workbook(data)
    assert summary["direct"] == 1
    assert summary["unresolved_causes"].get("unknown_distributor") == 1


def test_a_file_with_no_resolvable_rows_still_fails_loudly():
    """The reason the original rule existed: a workbook that is entirely from an
    unconfigured distributor should say so, not return a sheet full of blanks."""
    seed_price_list(_mh_list())
    data = seed_usage([{"filename": "ticket-1", "entity": "Someone Else Entirely",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    with pytest.raises(EnrichmentError) as exc:
        enrich_workbook(data)
    assert exc.value.reason == "unconfigured_distributor"


def _two_tabs(mh_price=None, mo_price=None, mh_hosp="Blake Hospital (HCA)",
              mo_hosp="Blake Hospital (HCA)"):
    tabs = {}
    if mh_price is not None:
        tabs[MH_TAB] = {"meta": ["Item", "Description"], "hospitals": [mh_hosp],
                        "rows": [("MTUUX***-GK", ["TIBIAL BASE PLATE (TITAN)"],
                                  {mh_hosp: mh_price})]}
    if mo_price is not None:
        tabs[MO_TAB] = {"meta": ["Item", "Class", "Part Type"], "hospitals": [mo_hosp],
                        "rows": [("MTUUX***-GK", ["Knee", "Tibial"],
                                  {mo_hosp: mo_price})]}
    return tabs


def test_a_hospital_only_on_another_tab_is_priced_as_an_estimate():
    """Blake Medical Center and Parkridge are priced only on 'MH for MO', and
    the real file's tickets are all Maxx Orthopedics. Those 15 rows came back
    blank before. They are priced now — but rose, because it is another sales
    channel's number for the account."""
    seed_price_list(_two_tabs(mh_price=925))       # nothing on the MO tab
    data = seed_usage([{"filename": "MO18711-A", "entity": "Maxx Orthopedics",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    out, summary = enrich_workbook(data)
    assert summary["estimates"] == 1 and summary["direct"] == 0
    assert summary["off_tab"] == ["Blake Hospital -> MH for MO"]
    _, value, rgb = usage_prices(out)[0]
    assert value == 925 and rgb.endswith("FFC7CE")


def test_the_own_tab_still_wins_when_both_tabs_have_the_hospital():
    seed_price_list(_two_tabs(mh_price=5555, mo_price=925))
    data = seed_usage([{"filename": "MO18711-A", "entity": "Maxx Orthopedics",
                        "hospital": "Blake Hospital", "ref": "MTUUX100-GK"}])
    out, summary = enrich_workbook(data)
    assert summary["direct"] == 1 and summary["off_tab"] == []
    _, value, rgb = usage_prices(out)[0]
    assert value == 925 and rgb.endswith("39FF14")


def test_a_better_match_on_another_tab_beats_a_fuzzy_one_at_home():
    """The worst case in the real file. 'Methodist Hospital HCA' matched its own
    tab only fuzzily — to 'Methodist Hospital Southlake', a different facility —
    while the other tab held 'Methodist Hospital (HCA)' exactly. Preferring the
    home tab regardless would have put a confidently wrong number in 21 cells."""
    seed_price_list(_two_tabs(mh_price=925, mo_price=5555,
                              mh_hosp="Methodist Hospital (HCA)",
                              mo_hosp="Methodist Hospital Southlake"))
    data = seed_usage([{"filename": "MO18711-A", "entity": "Maxx Orthopedics",
                        "hospital": "Methodist Hospital HCA", "ref": "MTUUX100-GK"}])
    out, summary = enrich_workbook(data)
    assert usage_prices(out)[0][1] == 925, "took the Southlake price"
    assert summary["off_tab"] == ["Methodist Hospital HCA -> MH for MO"]


def test_bracketed_and_unbracketed_system_markers_match():
    """The price list writes 'X (HCA)', the usage sheet writes 'X HCA'.
    normalize_hospital strips the parenthetical — which is what makes
    'Centerpoint Med Ctr (HCA)' match 'Centerpoint Medical Center' — but that
    same strip destroys this match instead of making it."""
    m = mt.match_hospital("Methodist Hospital HCA",
                          ["Methodist Hospital Southlake", "Methodist Hospital (HCA)"])
    assert m.method == "exact" and m.name == "Methodist Hospital (HCA)"


def test_stripping_still_matches_the_centerpoint_case():
    """The form that motivated stripping in the first place must not regress."""
    m = mt.match_hospital("Centerpoint Medical Center", ["Centerpoint Med Ctr (HCA)"])
    assert m.method == "exact"
