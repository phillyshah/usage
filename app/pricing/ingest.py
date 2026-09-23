"""Parse the hospital price-list workbook into (tab, item, hospital, price) rows.

Not built on ``ingest_reference._records`` on purpose: that helper returns rows
keyed by header name, and these tabs have up to 178 hospital columns including
literal duplicates ('Surgcenter of Plano' appears twice on the Summary tab), so
header-keyed dicts would silently collapse columns. This reads positionally.

The tabs also disagree about how many leading metadata columns they have —
'MH for MO' is Item/Description, 'Summary Price List' is Item/Class/Part Type —
so the first hospital column is detected, never assumed.
"""
from __future__ import annotations

import io
import logging

from openpyxl import load_workbook

from app.learning.ingest_reference import _cell_str, _is_xlsx
from app.pipeline.assemble import _money
from app.pricing.tabs import META_HEADERS

log = logging.getLogger("pricing.ingest")

MAX_HEADER_SCAN = 8


def _find_header_row(grid: list[list]) -> int | None:
    """Index of the header row: the first whose leading cell reads 'Item', else
    the first with at least three non-empty cells."""
    for i, row in enumerate(grid[:MAX_HEADER_SCAN]):
        first = (_cell_str(row[0]) or "").strip().lower() if row else ""
        if first == "item":
            return i
    for i, row in enumerate(grid[:MAX_HEADER_SCAN]):
        if sum(1 for c in row if _cell_str(c)) >= 3:
            return i
    return None


def _first_hospital_col(header: list) -> int:
    """Index of the first hospital column — one past the last metadata header."""
    last_meta = -1
    for i, cell in enumerate(header):
        label = (_cell_str(cell) or "").strip().lower()
        if label in META_HEADERS:
            last_meta = i
    return last_meta + 1 if last_meta >= 0 else 1


def parse_price_list(data: bytes) -> dict:
    """bytes -> {"rows": [...], "tabs": {tab: {"items": n, "hospitals": n, "prices": n}}}.

    Rows are dicts ready for ``db.replace_hospital_prices``.
    """
    if not _is_xlsx(data):
        raise ValueError("The price list must be an Excel workbook (.xlsx).")

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows: list[dict] = []
    tabs: dict[str, dict] = {}
    try:
        for tab in wb.sheetnames:
            grid = [list(r) for r in wb[tab].iter_rows(values_only=True)]
            hdr_i = _find_header_row(grid)
            if hdr_i is None:
                log.warning("price list tab %r has no header row; skipped", tab)
                continue
            header = grid[hdr_i]
            start = _first_hospital_col(header)
            hospitals = {
                i: name for i in range(start, len(header))
                if (name := (_cell_str(header[i]) or "").strip())
            }
            if not hospitals:
                log.warning("price list tab %r has no hospital columns; skipped", tab)
                continue

            # (item, hospital) -> price, so a duplicated column can be reconciled
            # before it reaches the table's composite primary key.
            seen: dict[tuple[str, str], float | None] = {}
            descriptions: dict[str, str | None] = {}
            for grow in grid[hdr_i + 1:]:
                if not grow:
                    continue
                code = (_cell_str(grow[0]) or "").strip()
                if not code:
                    continue
                desc = _cell_str(grow[start - 1]) if start >= 1 and len(grow) >= start else None
                descriptions.setdefault(code, desc)
                for i, hosp in hospitals.items():
                    price = _money(grow[i]) if i < len(grow) else None
                    if price is None:
                        continue
                    key = (code, hosp)
                    if key in seen and seen[key] != price:
                        # The same hospital listed twice at different prices is
                        # genuinely ambiguous — drop it rather than pick one.
                        log.warning("price list %r: %s @ %s has conflicting prices "
                                    "(%s vs %s); dropped", tab, code, hosp,
                                    seen[key], price)
                        seen[key] = None
                        continue
                    seen.setdefault(key, price)

            n = 0
            for (code, hosp), price in seen.items():
                if price is None:
                    continue
                rows.append({
                    "tab": tab, "item_code": code,
                    "description": descriptions.get(code),
                    "hospital": hosp, "unit_price": float(price),
                })
                n += 1
            tabs[tab] = {"items": len(descriptions), "hospitals": len(hospitals),
                         "prices": n}
    finally:
        wb.close()

    if not rows:
        raise ValueError("No prices found in the workbook — check it is the "
                         "hospital price list and that the tabs are intact.")
    return {"rows": rows, "tabs": tabs}


def ingest_price_list(data: bytes) -> dict:
    """Parse + full-replace. Returns a flat summary for the route/UI."""
    from app.db import db

    parsed = parse_price_list(data)
    db.replace_hospital_prices(parsed["rows"])
    return {
        "tabs": len(parsed["tabs"]),
        "prices": len(parsed["rows"]),
        "hospitals": sum(t["hospitals"] for t in parsed["tabs"].values()),
        "per_tab": parsed["tabs"],
    }
