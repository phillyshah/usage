"""Step 5 — fill the blank Price cells of a generated usage workbook.

A blank Price on the Usage sheet means the ticket quoted a *construct* price
rather than a per-component one, so the accountant has been typing those in by
hand against the hospital price list. This does the lookup and hands the same
workbook back with those cells filled:

  * neon green (39FF14) — the price came straight out of the price list for
    this hospital and this component. Nothing to check.
  * rose (FFC7CE) — an inferred estimate. Defensible starting number; review it.
  * still red — no evidence at all. Unchanged, exactly as it arrived.

Everything else in the file is left alone, and that is asserted rather than
assumed: the workbook is snapshotted before and after, and the run fails without
publishing if any cell outside the planned set moved.
"""
from __future__ import annotations

import io
import logging
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles import PatternFill

from app.db import db
from app.pipeline.template import MAXX_HEALTH, MAXX_ORTHO, detect_template
from app.pricing import match as mt
from app.pricing import normalize as nz
from app.pricing.estimate import UsageRow, estimate_price
from app.pricing.tabs import DISTRIBUTOR_TABS, LEGAL_SUFFIXES

log = logging.getLogger("pricing.enrich")

GREEN = PatternFill(start_color="FF39FF14", end_color="FF39FF14", fill_type="solid")
ROSE = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
YELLOW_RGB = "FFFF00"

# Workbook parts openpyxl cannot round-trip. Our own generated workbooks contain
# none of these, so the normal path never trips it — but a file the operator has
# decorated in Excel would silently lose them on save, and silently destroying
# someone's work is worse than refusing to touch it.
LOSSY_PARTS = ("xl/charts/", "xl/media/", "xl/drawings/", "vbaProject.bin")

# The run id is stamped into the workbook's custom document properties rather
# than a cell: it lives outside the grid, so it disturbs nothing the operator
# sees and the "only blank Price cells changed" validator never sees it either.
# When the file comes back through step 4 it is what says which run to diff
# against — without it a returned workbook is just a column of numbers with no
# way to tell a human's decision from the tool's guess.
RUN_ID_PROP = "UsagePricingRunId"


def stamped_run_id(wb) -> str | None:
    try:
        for prop in wb.custom_doc_props.props:
            if prop.name == RUN_ID_PROP:
                return str(prop.value) or None
    except Exception:  # pragma: no cover - absent or unreadable props
        pass
    return None


class EnrichmentError(ValueError):
    """A run that must fail without publishing. The message is operator-facing."""

    def __init__(self, message: str, reason: str):
        super().__init__(message)
        self.reason = reason


@dataclass
class CellPlan:
    excel_row: int
    value: float
    kind: str                 # direct | estimate
    basis: str
    tab: str
    hospital: str | None
    codes: tuple[str, ...] = ()


@dataclass
class RunSummary:
    tabs: list[str] = field(default_factory=list)
    eligible: int = 0
    direct: int = 0
    estimates: int = 0
    unresolved: int = 0
    skipped_wasted: int = 0
    zero_estimates: int = 0
    # Why the unresolved rows are unresolved. Without this the first real run
    # looks broken: most blanks are blank because the Hospital or Ref cell
    # upstream of them is itself blank, which pricing can do nothing about.
    unresolved_causes: dict[str, int] = field(default_factory=dict)
    fuzzy_hospitals: list[str] = field(default_factory=list)
    ambiguous_hospitals: list[str] = field(default_factory=list)
    # Rows priced from a tab other than the ticket's own distributor's. Always
    # written rose, and listed here so it is visible which accounts are being
    # priced from another sales channel's list.
    off_tab: list[str] = field(default_factory=list)

    def note(self, cause: str) -> None:
        self.unresolved += 1
        self.unresolved_causes[cause] = self.unresolved_causes.get(cause, 0) + 1

    def as_dict(self) -> dict:
        return {
            "tabs": self.tabs, "eligible": self.eligible, "direct": self.direct,
            "estimates": self.estimates, "unresolved": self.unresolved,
            "skipped_wasted": self.skipped_wasted,
            "zero_estimates": self.zero_estimates,
            "unresolved_causes": self.unresolved_causes,
            "fuzzy_hospitals": sorted(set(self.fuzzy_hospitals)),
            "ambiguous_hospitals": sorted(set(self.ambiguous_hospitals)),
            "off_tab": sorted(set(self.off_tab)),
        }


# --------------------------------------------------------------------------
# Workbook inspection
# --------------------------------------------------------------------------
def _guard_lossy_content(data: bytes) -> None:
    try:
        names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    except zipfile.BadZipFile:
        raise EnrichmentError(
            "That file isn't a readable .xlsx workbook.", "parse_failure") from None
    found = sorted({p.rstrip("/").rsplit("/", 1)[-1] or p
                    for p in LOSSY_PARTS
                    if any(n.startswith(p) or n.endswith(p) for n in names)})
    if found:
        raise EnrichmentError(
            "This workbook contains charts, images or macros, which cannot be "
            "preserved when prices are written back. Upload the workbook as "
            "generated in step 3, or strip the extras first.",
            "unsupported_workbook_content")


def _headers(ws) -> dict[str, int]:
    return {str(c.value).strip(): i for i, c in enumerate(ws[1], start=1)
            if c.value is not None}


def _fill_rgb(cell) -> str | None:
    f = cell.fill
    if f is None or f.fill_type is None:
        return None
    rgb = getattr(f.start_color, "rgb", None)
    if isinstance(rgb, str):
        return rgb
    return f"theme:{getattr(f.start_color, 'theme', None)}"


def is_truly_blank(cell) -> bool:
    """Blank enough to fill.

    A formula is never blank even when it displays empty — overwriting ``=""``
    would destroy the operator's own work. ``0`` is a real price of zero and is
    not a hole. Only None and whitespace qualify.
    """
    if cell.data_type == "f":
        return False
    v = cell.value
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() == "" and not v.strip().startswith("=")
    return False


def _number_format(ws, col: int, row: int, span: int = 50) -> str:
    """The currency style of the nearest formatted Price cell, so an inserted
    price looks like the ones already in the column."""
    for r in list(range(row - 1, max(1, row - span) - 1, -1)) + \
             list(range(row + 1, min(ws.max_row, row + span) + 1)):
        fmt = ws.cell(row=r, column=col).number_format
        if fmt and fmt != "General":
            return fmt
    return "General"


# --------------------------------------------------------------------------
# Entity -> tab
# --------------------------------------------------------------------------
def _entity_by_image(wb) -> dict[str, str]:
    """Source Image stem -> Entity, read off the Tickets sheet."""
    if "Tickets" not in wb.sheetnames:
        return {}
    ws = wb["Tickets"]
    idx = _headers(ws)
    img_c, ent_c = idx.get("Source Image"), idx.get("Entity")
    if not img_c or not ent_c:
        return {}
    out: dict[str, str] = {}
    for r in range(2, ws.max_row + 1):
        stem = ws.cell(row=r, column=img_c).value
        ent = ws.cell(row=r, column=ent_c).value
        if stem and ent:
            out[str(stem).strip()] = str(ent).strip()
    return out


def normalize_entity(text: str | None) -> str:
    """The form ``DISTRIBUTOR_TABS`` keys are written in.

    ``Entity`` is the model's reading of printed text, so one batch of tickets
    from one company arrives spelled several ways — the first real workbook had
    ``Maxx Orthopedics``, ``Maxx Orthopedics, Inc`` and ``Maxx Orthopedics, Inc.``
    across 37 tickets. Punctuation and legal suffixes are dropped, reusing the
    suffix list that already exists for hospitals.
    """
    if not text:
        return ""
    toks = [t for t in nz._NON_ALNUM.sub(" ", str(text).lower()).split()
            if t not in LEGAL_SUFFIXES]
    return " ".join(toks)


def resolve_entity(stem: str | None, entity_map: dict[str, str]) -> tuple[str, str]:
    """``(entity, preferred_tab)`` for one Usage row, or ``("", "")``.

    **The filename is tried first, and that ordering is the fix.** ``Entity`` is
    a vision read of printed text and varies; the ``MH``/``MO`` filename prefix
    is a naming convention the business actually follows, and ``detect_template``
    already treats it as authoritative for PHI redaction. Preferring the shakier
    signal is what made a whole 153-row workbook fail on one ticket that read
    ``MAXX`` — a string no amount of normalizing can resolve, while its filename
    ``MO18711-A`` says Maxx Orthopedics outright.

    Returns empty strings rather than raising: one unresolvable row should cost
    that row, not the run. The caller fails the run only if nothing resolves.
    """
    stem = (stem or "").strip()
    entity = ""
    if stem:
        guessed = detect_template(None, stem)
        if guessed in (MAXX_HEALTH, MAXX_ORTHO):
            entity = guessed
    written = (entity_map.get(stem) or "").strip()
    if not entity:
        entity = written
    elif written and normalize_entity(written) != normalize_entity(entity):
        log.info("row image %s: filename says %r, ticket says %r — using the "
                 "filename", stem, entity, written)
    tab = DISTRIBUTOR_TABS.get(normalize_entity(entity)) if entity else ""
    return entity, tab or ""


# --------------------------------------------------------------------------
# Per-tab index
# --------------------------------------------------------------------------
@dataclass
class TabIndex:
    tab: str
    hospitals: list[str]
    catalog: dict[str, str]                      # item_code -> description
    prices: dict[str, dict[str, float]]          # item_code -> {hospital: price}


# Hospital-match quality, best first. A tab that knows the hospital exactly is
# a better source than one that only recognises it fuzzily, whichever tab the
# ticket's own distributor uses — see choose_tab.
_TIER_RANK = {"exact": 0, "alias": 1, "core": 2, "fuzzy": 3}


def load_all_tabs() -> dict[str, TabIndex]:
    """Every tab in the stored price list, not only the configured ones."""
    return {tab: load_tab(tab) for tab in db.hospital_price_tabs()}


def choose_tab(hospital: str | None, preferred: str,
               indexes: dict[str, TabIndex]) -> tuple[str, "mt.HospitalMatch"] | None:
    """Which tab to price this hospital from, and how well it matched.

    The tabs turned out to be *where an account is listed*, not a price
    agreement scoped to one distributor: on the first real workbook, 27 of 88
    blank prices belonged to hospitals that are priced only on a tab the old
    one-tab-per-entity rule forbade. So every tab is searched.

    Ranked by match quality first and only then by the ticket's own tab, because
    preference alone picks the wrong answer where it matters most: on that file
    'Methodist Hospital HCA' matched its own tab fuzzily (to 'Methodist Hospital
    Southlake', a different facility) while another tab had it exactly. Quality
    first fixes those 21 rows; the tie-break still keeps a row at home whenever
    two tabs know the hospital equally well.
    """
    scored = []
    for tab, index in indexes.items():
        m = mt.match_hospital(hospital, index.hospitals)
        if m.matched:
            scored.append((_TIER_RANK.get(m.method, 9), tab != preferred, tab, m))
    if not scored:
        return None
    scored.sort(key=lambda t: (t[0], t[1], t[2]))
    best_tier = scored[0][0]
    tied = [t for t in scored if t[0] == best_tier]
    # Two different tabs knowing the hospital equally well, and neither of them
    # the ticket's own, is a genuine coin flip. Say so instead of picking.
    if len(tied) > 1 and all(tab != preferred for _, _, tab, _ in tied):
        return None
    return scored[0][2], scored[0][3]


def load_tab(tab: str) -> TabIndex:
    rows = db.hospital_prices_for_tab(tab)
    if not rows:
        raise EnrichmentError(
            f'The price list has no tab named "{tab}". Upload the hospital '
            "price list on the Reference Data tile first.", "missing_tab")
    hospitals: dict[str, None] = {}
    catalog: dict[str, str] = {}
    prices: dict[str, dict[str, float]] = {}
    for r in rows:
        code, hosp = r["item_code"], r["hospital"]
        hospitals.setdefault(hosp, None)
        if r.get("description"):
            catalog.setdefault(code, r["description"])
        else:
            catalog.setdefault(code, "")
        prices.setdefault(code, {})[hosp] = float(r["unit_price"])
    return TabIndex(tab, list(hospitals), catalog, prices)


# --------------------------------------------------------------------------
# Pricing one row
# --------------------------------------------------------------------------
def _direct_price(hosp: mt.HospitalMatch, comp: mt.ComponentMatch,
                  index: TabIndex) -> float | None:
    """One unambiguous price-list value, or None.

    Only the winning component tier contributes — see match.py for why that
    short-circuit matters. Values are rounded to cents before comparison so a
    cached 1200.0000000001 doesn't read as a disagreement.
    """
    if not (hosp.confident and comp.matched):
        return None
    vals = {round(index.prices.get(c, {})[hosp.name], 2)
            for c in comp.codes if hosp.name in index.prices.get(c, {})}
    if len(vals) != 1:
        return None
    return vals.pop()


def price_row(row: UsageRow, preferred_tab: str, indexes: dict[str, TabIndex],
              observations: list[UsageRow], summary: RunSummary) -> CellPlan | None:
    if not row.hospital:
        summary.note("no_hospital")
        return None
    if not preferred_tab:
        summary.note("unknown_distributor")
        return None

    chosen = choose_tab(row.hospital, preferred_tab, indexes)
    if chosen is None:
        # No tab knows this hospital (or two knew it equally well and neither
        # was the ticket's own). That rules out a price-list lookup, but NOT an
        # estimate: rungs 1, 2, 4 and 6 are built from other rows of this same
        # workbook and never touch the price list. Carry on with the ticket's
        # own tab and an unmatched hospital.
        tab, hosp = preferred_tab, mt.HospitalMatch(None, "none")
    else:
        tab, hosp = chosen
    index = indexes[tab]
    at_home = tab == preferred_tab

    if hosp.method == "fuzzy":
        summary.fuzzy_hospitals.append(f"{row.hospital} -> {hosp.name}")
    if not at_home:
        summary.off_tab.append(f"{row.hospital} -> {tab}")

    comp = mt.match_component(row.ref, row.description, index.catalog)

    direct = _direct_price(hosp, comp, index)
    if direct is not None:
        if at_home:
            return CellPlan(row.excel_row, direct, "direct",
                            f"price list: {hosp.name}", tab, hosp.name,
                            comp.codes)
        # The price is real, but it is another sales channel's number for this
        # account. Green says "came straight from your list, don't check it",
        # which is a promise only the ticket's own tab can make.
        return CellPlan(row.excel_row, direct, "estimate",
                        f"price list: {hosp.name} (on the {tab} tab)",
                        tab, hosp.name, comp.codes)

    est = estimate_price(row, observations, index.prices, index.catalog,
                         comp.codes, comp.fallbacks, hosp.name)
    if est.found:
        if est.value == 0 and not est.zero_flagged:
            summary.note("zero_without_direct_evidence")
            return None
        if est.value == 0:
            summary.zero_estimates += 1
        return CellPlan(row.excel_row, float(est.value), "estimate", est.basis,
                        tab, hosp.name, comp.codes)

    if not hosp.matched:
        summary.note("hospital_not_in_price_list")
    elif not row.ref:
        summary.note("no_ref")
    elif not comp.matched:
        summary.note("component_not_in_price_list")
    else:
        summary.note("no_supporting_evidence")
    return None


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------
def _snapshot(wb) -> dict[tuple[str, int, int], tuple]:
    snap: dict[tuple[str, int, int], tuple] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                snap[(name, cell.row, cell.column)] = (
                    cell.value, cell.data_type, cell.number_format,
                    _fill_rgb(cell))
    return snap


def _validate(before: dict, after: dict, plans: list[CellPlan],
              price_col: int) -> None:
    changed = {k for k in set(before) | set(after)
               if before.get(k) != after.get(k)}
    planned = {("Usage", p.excel_row, price_col) for p in plans}
    if changed != planned:
        stray = sorted(changed - planned)[:5]
        missed = sorted(planned - changed)[:5]
        raise EnrichmentError(
            "Internal check failed: the workbook changed in places this step "
            f"did not plan to touch (unexpected={stray}, missing={missed}). "
            "Nothing was published.", "validation_failed")
    by_row = {p.excel_row: p for p in plans}
    for name, r, c in changed:
        if not (before.get((name, r, c), (None,))[0] is None
                or str(before[(name, r, c)][0]).strip() == ""):
            raise EnrichmentError(
                f"Internal check failed: row {r} was not blank before. "
                "Nothing was published.", "validation_failed")
        want = "39FF14" if by_row[r].kind == "direct" else "FFC7CE"
        got = after[(name, r, c)]
        if not (got[3] or "").endswith(want):
            raise EnrichmentError(
                f"Internal check failed: row {r} was not coloured for a "
                f"{by_row[r].kind}. Nothing was published.", "validation_failed")
        if not isinstance(got[0], (int, float)):
            raise EnrichmentError(
                f"Internal check failed: row {r} was filled with a non-number. "
                "Nothing was published.", "validation_failed")


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------
def enrich_workbook(data: bytes, run_id: str | None = None) -> tuple[bytes, dict, list[dict]]:
    """Usage workbook bytes -> (enriched bytes, summary, cells).

    ``cells`` is what was written, one record per filled cell: enough for step 4
    to tell a reviewed number from an untouched one. Keyed on ref + hospital
    because that is what ``learn_price`` needs — no line id required, which is
    just as well, since Usage and Line Items are not row-for-row (153 vs 159 on
    the first real file).

    Raises EnrichmentError on anything that must not publish.
    """
    _guard_lossy_content(data)

    # data_only=False (the default, stated for the record): data_only=True would
    # replace every formula in the file with its cached value on save.
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False)
    except Exception as exc:
        raise EnrichmentError(f"Could not read that workbook: {exc}",
                              "parse_failure") from exc

    if "Usage" not in wb.sheetnames:
        raise EnrichmentError(
            'That workbook has no "Usage" sheet. Upload the workbook the app '
            "generated in step 3.", "missing_usage_columns")
    ws = wb["Usage"]
    idx = _headers(ws)
    required = ["Source Image Filename", "Hospital", "Price", "Ref Number"]
    missing = [h for h in required if h not in idx]
    if missing:
        raise EnrichmentError(
            f"The Usage sheet is missing required column(s): {', '.join(missing)}.",
            "missing_usage_columns")
    price_col = idx["Price"]
    merged = {c for rng in ws.merged_cells.ranges for c in rng.cells}

    summary = RunSummary()
    entity_map = _entity_by_image(wb)
    indexes = load_all_tabs()
    if not indexes:
        raise EnrichmentError(
            "The price list has not been uploaded yet. Add it on the Reference "
            "Data tile before running step 5.", "missing_tab")
    rows_by_entity: dict[str, list[UsageRow]] = {}
    by_row_ref: dict[int, str | None] = {}
    blanks: list[tuple[UsageRow, str, str]] = []
    resolved_rows = 0

    for r in range(2, ws.max_row + 1):
        stem = ws.cell(row=r, column=idx["Source Image Filename"]).value
        if stem is None and all(ws.cell(row=r, column=c).value is None
                                for c in idx.values()):
            continue
        entity, tab = resolve_entity(
            str(stem) if stem is not None else None, entity_map)
        if tab:
            resolved_rows += 1
            if tab not in summary.tabs:
                summary.tabs.append(tab)

        cell = ws.cell(row=r, column=price_col)
        ref = ws.cell(row=r, column=idx["Ref Number"]).value
        info = db.part_info_for_ref(str(ref).strip()) if ref else None
        urow = UsageRow(
            excel_row=r,
            hospital=(lambda v: str(v).strip() if v else None)(
                ws.cell(row=r, column=idx["Hospital"]).value),
            ref=str(ref).strip() if ref else None,
            description=(info or {}).get("description"),
            part_type=(info or {}).get("part_type"),
            category=(info or {}).get("category"),
            price=cell.value if isinstance(cell.value, (int, float)) else None,
        )
        # Observations are partitioned by DISTRIBUTOR, not by tab: "same part
        # type elsewhere in this file" must not import a Maxx Orthopedics price
        # into a Maxx Health row, which a mixed batch makes possible. Tabs are
        # no longer distributor-scoped, so the entity is what this keys on.
        rows_by_entity.setdefault(normalize_entity(entity), []).append(urow)
        by_row_ref[r] = urow.ref

        if (r, price_col) in merged:
            continue
        if not is_truly_blank(cell):
            continue
        if (_fill_rgb(cell) or "").endswith(YELLOW_RGB):
            # Wasted component. A wasted line's price is a business decision,
            # not a lookup — leave it yellow and say so in the summary.
            summary.skipped_wasted += 1
            continue
        summary.eligible += 1
        blanks.append((urow, tab, normalize_entity(entity)))

    # A row whose distributor cannot be identified costs that row, not the run —
    # but a file where NOTHING resolves is a genuinely unconfigured distributor
    # and should say so loudly rather than return a workbook full of blanks.
    if blanks and resolved_rows == 0:
        raise EnrichmentError(
            "No price-list tab is configured for any ticket in this file. Check "
            "the Entity column, or add the distributor to DISTRIBUTOR_TABS.",
            "unconfigured_distributor")

    plans: list[CellPlan] = []
    for urow, tab, ent_key in blanks:
        plan = price_row(urow, tab, indexes, rows_by_entity.get(ent_key, []),
                         summary)
        if plan is not None:
            plans.append(plan)

    before = _snapshot(wb)
    for p in plans:
        cell = ws.cell(row=p.excel_row, column=price_col)
        cell.value = round(float(p.value), 2)
        cell.number_format = _number_format(ws, price_col, p.excel_row)
        cell.fill = GREEN if p.kind == "direct" else ROSE
        if p.kind == "direct":
            summary.direct += 1
        else:
            summary.estimates += 1
    _validate(before, _snapshot(wb), plans, price_col)

    cells = [{"row": p.excel_row, "ref": by_row_ref.get(p.excel_row),
              "hospital": p.hospital, "value": round(float(p.value), 2),
              "kind": p.kind, "tab": p.tab}
             for p in plans]
    if run_id:
        wb.custom_doc_props.append(StringProperty(name=RUN_ID_PROP, value=run_id))

    if summary.eligible != summary.direct + summary.estimates + summary.unresolved:
        raise EnrichmentError(
            "Internal check failed: the run summary does not add up "
            f"({summary.eligible} eligible vs {summary.direct} direct + "
            f"{summary.estimates} estimated + {summary.unresolved} unresolved). "
            "Nothing was published.", "validation_failed")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), summary.as_dict(), cells
